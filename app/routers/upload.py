"""
MCPilot — Document Upload Router
Accepts PDF, DOCX, TXT uploads. Extracts text server-side with
OCR fallback (embedded image extraction + Tesseract) for scanned PDFs.
Handles ICCBased colorspace common in Indian hospital documents.

Endpoints:
  POST /gateway/upload/analyse   — upload + tool → result
  POST /gateway/upload/extract   — upload → extracted text only
"""
import asyncio
import io
from concurrent.futures import ThreadPoolExecutor
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from app.core.logging import get_logger
from app.compliance.classification import registry as clearance_registry
from app.compliance.pipeline import scan_output

logger = get_logger(__name__)
router = APIRouter(prefix="/gateway/upload", tags=["Document Upload"])

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "doc",
    "text/plain": "txt",
    "text/csv":   "csv",
}
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB
MAX_PAGES     = 50
_executor     = ThreadPoolExecutor(max_workers=2)


def _ocr_pdf(content: bytes) -> str:
    """
    OCR scanned PDF using PyMuPDF + Tesseract.
    Uses fitz.Pixmap for colorspace conversion — handles ICCBased JPEG
    common in Indian hospital scanned documents.
    """
    import fitz
    import pytesseract
    from PIL import Image

    text_parts = []
    doc   = fitz.open(stream=content, filetype="pdf")
    pages = min(len(doc), MAX_PAGES)
    logger.info(f"OCR: processing {pages} pages")

    for page_num in range(pages):
        page = doc[page_num]
        text = ""

        # Try embedded image via fitz.Pixmap (handles ICC colorspace)
        images = page.get_images(full=True)
        if images:
            try:
                xref = images[0][0]
                pix  = fitz.Pixmap(doc, xref)
                # Convert to RGB — handles CMYK, ICCBased, Grayscale
                if pix.colorspace and pix.colorspace.n != 3:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                img  = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img, lang="eng")
            except Exception as e:
                logger.warning(f"Embedded image OCR failed p{page_num+1}: {e}")

        # Fallback: render full page at 200 DPI
        if not text.strip():
            try:
                mat = fitz.Matrix(200/72, 200/72)
                pix = page.get_pixmap(matrix=mat)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img, lang="eng")
            except Exception as e:
                logger.warning(f"Render OCR failed p{page_num+1}: {e}")

        if text.strip():
            text_parts.append(f"[Page {page_num+1}]\n{text.strip()}")

    doc.close()
    result = "\n\n".join(text_parts).strip()
    if not result:
        raise ValueError(
            "OCR could not extract text. Document may be encrypted or corrupted."
        )
    return result


def _extract_text_from_pdf(content: bytes) -> str:
    """
    Extract text from PDF.
    1. pdfplumber for text-based PDFs (fast, <1s)
    2. OCR fallback for scanned/image PDFs (15-30s)
    """
    import pdfplumber

    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages[:MAX_PAGES]:
                text = page.extract_text()
                if text and text.strip():
                    text_parts.append(text)
    except Exception as e:
        logger.warning(f"pdfplumber failed: {e}")

    combined = "\n\n".join(text_parts).strip()
    if len(combined) > 100:
        logger.info(f"Text PDF: {len(combined.split())} words")
        return combined

    # Scanned — use OCR
    logger.info("No text found — switching to OCR")
    return _ocr_pdf(content)


def _extract_text_from_docx(content: bytes) -> str:
    import docx
    doc  = docx.Document(io.BytesIO(content))
    rows = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            r = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if r: rows.append(r)
    return "\n".join(rows)


def _extract_text_from_txt(content: bytes) -> str:
    try:    return content.decode("utf-8")
    except: return content.decode("latin-1")


def extract_text(content: bytes, file_type: str) -> str:
    if file_type == "pdf":              return _extract_text_from_pdf(content)
    elif file_type in ("docx", "doc"): return _extract_text_from_docx(content)
    elif file_type in ("txt", "csv"):  return _extract_text_from_txt(content)
    else: raise ValueError(f"Unsupported type: {file_type}")


def _detect_type(content_type: str, filename: str) -> str | None:
    ft = ALLOWED_TYPES.get(content_type)
    if not ft and filename:
        ext = filename.rsplit(".", 1)[-1].lower()
        if ext in ("pdf", "docx", "doc", "txt", "csv"):
            ft = ext
    return ft


# ══════════════════════════════════════════════════════════════════════════════
# ENDPOINT: /gateway/upload/extract
# ══════════════════════════════════════════════════════════════════════════════
@router.post("/extract", summary="Extract text from uploaded document")
async def extract_document_text(
    request: Request,
    file: UploadFile = File(...),
) -> dict:
    """Extract plain text from PDF, DOCX or TXT. OCR fallback for scanned PDFs."""
    ft = _detect_type(file.content_type or "", file.filename or "")
    if not ft:
        raise HTTPException(status_code=415, detail="Unsupported file type.")
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413,
            detail=f"File too large. Max {MAX_FILE_SIZE//1024//1024}MB.")
    if not content:
        raise HTTPException(status_code=400, detail="Empty file.")
    loop = asyncio.get_event_loop()
    try:
        text = await loop.run_in_executor(_executor, extract_text, content, ft)
    except Exception as e:
        raise HTTPException(status_code=422, detail=str(e))
    return {
        "status":       "OK",
        "filename":     file.filename,
        "file_type":    ft.upper(),
        "ocr_used":     "[Page 1]" in text,
        "word_count":   len(text.split()),
        "text_preview": text[:500] + "..." if len(text) > 500 else text,
        "full_text":    text,
    }


# ══════════════════════════════════════════════════════════════════════════════
# ENDPOINT: /gateway/upload/analyse
# ══════════════════════════════════════════════════════════════════════════════
@router.post("/analyse", summary="Upload document and analyse with MCPilot tool")
async def analyse_document(
    request:      Request,
    file:         UploadFile = File(...),
    server_id:    str        = Form(...),
    tool_name:    str        = Form(...),
    extra_params: str        = Form("{}"),
) -> dict:
    """
    Upload document and analyse with MCPilot tool.
    OCR handles scanned Indian hospital documents automatically.
    """
    import json

    # Classification gate
    tenant_id = getattr(request.state, "tenant_id", "unknown")
    gating    = clearance_registry.gate(requestor_tenant=tenant_id, target_server=server_id)
    if gating.denied:
        raise HTTPException(status_code=403, detail={
            "error":      "CLASSIFICATION_DENIED",
            "reason":     gating.reason,
            "clearance":  gating.requestor_clearance.name,
            "required":   gating.data_classification.name,
            "audit_hash": gating.audit_hash,
        })

    ft = _detect_type(file.content_type or "", file.filename or "")
    if not ft:
        raise HTTPException(status_code=415,
            detail="Unsupported file type. Use PDF, DOCX or TXT.")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413,
            detail=f"File too large. Max {MAX_FILE_SIZE//1024//1024}MB.")
    if not content:
        raise HTTPException(status_code=400, detail="Empty file.")

    # Extract text in thread pool (non-blocking, OCR-aware)
    loop = asyncio.get_event_loop()
    try:
        document_text = await loop.run_in_executor(_executor, extract_text, content, ft)
    except Exception as e:
        raise HTTPException(status_code=422, detail=str(e))

    if not document_text.strip():
        raise HTTPException(status_code=422, detail="No text extracted.")

    # Build params
    try:
        params = json.loads(extra_params) if extra_params and extra_params != "{}" else {}
    except json.JSONDecodeError:
        params = {}
    params["document_text"] = document_text

    # Call MCP tool
    manager = request.app.state.mcp_manager
    try:
        result = await manager.call_tool(
            server_id=server_id, tool_name=tool_name, parameters=params)
    except KeyError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Tool call failed: {e}")

    output_scan = scan_output(result, server_id=server_id)

    return {
        "status":       "ok",
        "server_id":    server_id,
        "tool_name":    tool_name,
        "filename":     file.filename,
        "file_type":    ft.upper(),
        "word_count":   len(document_text.split()),
        "ocr_used":     "[Page 1]" in document_text,
        "gating": {
            "action":     gating.action,
            "clearance":  gating.requestor_clearance.name,
            "audit_hash": gating.audit_hash,
        },
        "result":       output_scan.redacted,
        "pii_detected": output_scan.phi_detected,
    }
